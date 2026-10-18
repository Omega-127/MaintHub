from io import BytesIO
from datetime import date, timedelta
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


DEPARTMENT_NAMES = {
    "BLOWROOM":    "Blowroom",
    "COMBER":      "Comber",
    "RING_FRAME":  "Ring Frame",
    "SPEED_FRAME": "Speed Frame",
    "WINDING":     "Winding",
    "BUFFING":     "Buffing",
}


def format_date(dt):
    if not dt:
        return "-"
    if hasattr(dt, "strftime"):
        return dt.strftime("%d-%m-%Y")
    if isinstance(dt, str):
        try:
            from datetime import datetime
            return datetime.strptime(dt.split("T")[0], "%Y-%m-%d").strftime("%d-%m-%Y")
        except Exception:
            return dt
    return str(dt)


def build_department_sheet(ws, department, machines):
    """
    Builds a beautifully styled maintenance sheet for a department.
    Contains:
    - Title banner with the department name as title
    - Metadata subtitle (Generated on, Total machines)
    - Department, Machine Name, Type, Location, Interval, Status
    - Previous Maintenance Date and Due Date
    - Maintenance history columns (Previous Date & Due Date pairs for each historical cycle)
    - All machines in the department
    """
    today = date.today()
    dept_label = DEPARTMENT_NAMES.get(department, department)
    ws.title = dept_label[:31]
    ws.sheet_view.showGridLines = True

    # Compute max history length across machines in this department
    max_history = 0
    for m in machines:
        hist_count = len(m.maintenance_histories) if m.maintenance_histories else 0
        if hist_count > max_history:
            max_history = hist_count

    # Define headers
    base_headers = [
        "Sr. No.",
        "Department",
        "Machine Name",
        "Machine Type",
        "Location",
        "Interval (Days)",
        "Status",
        "Previous Maintenance Date",
        "Due Date",
    ]

    # Additional history columns (pairs of Previous Maintenance Date & Due Date)
    history_headers = []
    for i in range(1, max_history + 1):
        history_headers.append(f"Previous Maintenance Date (History {i})")
        history_headers.append(f"Due Date (History {i})")

    headers = base_headers + history_headers
    total_cols = len(headers)

    # Styles: 8-character ARGB hex strings (FF = 100% opaque alpha)
    # OpenPyXL pads 6-character hex strings with '00' which causes 100% transparent text in Excel viewers!
    font_title = Font(name="Calibri", size=15, bold=True, color="FFFFFFFF")
    font_sub   = Font(name="Calibri", size=10, italic=True, color="FF475569")
    font_hdr   = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
    font_data  = Font(name="Calibri", size=10, color="FF0F172A")
    font_bold  = Font(name="Calibri", size=10, bold=True, color="FF0F172A")

    fill_title = PatternFill(start_color="FF1E3A8A", end_color="FF1E3A8A", fill_type="solid")  # Navy Blue
    fill_sub   = PatternFill(start_color="FFF1F5F9", end_color="FFF1F5F9", fill_type="solid")
    fill_hdr   = PatternFill(start_color="FF2F6FE6", end_color="FF2F6FE6", fill_type="solid")  # Brand Blue
    fill_even  = PatternFill(start_color="FFFFFFFF", end_color="FFFFFFFF", fill_type="solid")
    fill_odd   = PatternFill(start_color="FFF8FAFC", end_color="FFF8FAFC", fill_type="solid")  # Slate-50

    align_center = Alignment(horizontal="center", vertical="center")
    align_left   = Alignment(horizontal="left",   vertical="center")
    align_right  = Alignment(horizontal="right",  vertical="center")
    align_hdr    = Alignment(horizontal="center", vertical="center", wrap_text=True)

    border_thin = Border(
        left=Side(style="thin", color="FFCBD5E1"),
        right=Side(style="thin", color="FFCBD5E1"),
        top=Side(style="thin", color="FFCBD5E1"),
        bottom=Side(style="thin", color="FFCBD5E1"),
    )

    # Row 1: Department Title
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    for col_c in range(1, total_cols + 1):
        c = ws.cell(row=1, column=col_c)
        c.fill = fill_title
        c.font = font_title
        c.alignment = align_center
    ws.cell(row=1, column=1, value=f"{dept_label.upper()} DEPARTMENT - MAINTENANCE SCHEDULE")
    ws.row_dimensions[1].height = 36

    # Row 2: Subtitle
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_cols)
    for col_c in range(1, total_cols + 1):
        c = ws.cell(row=2, column=col_c)
        c.fill = fill_sub
        c.font = font_sub
        c.alignment = align_center
    ws.cell(row=2, column=1, value=f"MaintHub Maintenance Management System  |  Generated on: {today.strftime('%d-%m-%Y')}  |  Total Machines: {len(machines)}")
    ws.row_dimensions[2].height = 20

    # Row 3: Blank separator
    ws.row_dimensions[3].height = 10

    # Row 4: Column Headers
    ws.row_dimensions[4].height = 28
    for col_idx, header_text in enumerate(headers, start=1):
        c = ws.cell(row=4, column=col_idx, value=header_text)
        c.font = font_hdr
        c.fill = fill_hdr
        c.alignment = align_hdr
        c.border = border_thin

    # Row 5+: Data rows for each machine
    current_row = 5
    for idx, m in enumerate(machines, start=1):
        ws.row_dimensions[current_row].height = 22
        row_fill = fill_even if idx % 2 != 0 else fill_odd

        # Calculate machine status display
        is_overdue = m.next_maintenance_date and m.next_maintenance_date < today and m.status == "ACTIVE"
        status_display = "OVERDUE" if is_overdue else (m.status or "ACTIVE")

        # Previous date & next due date for current cycle
        prev_date_str = format_date(m.last_maintenance_date)
        due_date_str  = format_date(m.next_maintenance_date)

        row_data = [
            idx,
            dept_label,
            m.name,
            m.type or "-",
            m.location or "-",
            m.maintenance_interval if m.maintenance_interval is not None else "-",
            status_display,
            prev_date_str,
            due_date_str,
        ]

        # Add history records chronologically (oldest to newest)
        histories = sorted(m.maintenance_histories, key=lambda h: h.maintenance_date) if m.maintenance_histories else []
        for h_idx in range(max_history):
            if h_idx < len(histories):
                h = histories[h_idx]
                h_prev = format_date(h.maintenance_date)
                # Compute scheduled due date for that history cycle
                if m.maintenance_interval:
                    h_due = format_date(h.maintenance_date.date() + timedelta(days=m.maintenance_interval))
                else:
                    h_due = "-"
                row_data.append(h_prev)
                row_data.append(h_due)
            else:
                row_data.append("-")
                row_data.append("-")

        for col_idx, val in enumerate(row_data, start=1):
            c = ws.cell(row=current_row, column=col_idx, value=val)
            c.font = font_data
            c.fill = row_fill
            c.border = border_thin

            # Specific alignments
            if col_idx in (1, 6):
                c.alignment = align_center
            elif col_idx in (2, 7, 8, 9) or col_idx > 9:
                c.alignment = align_center
            else:
                c.alignment = align_left

            # Highlight status
            if col_idx == 7:
                if status_display == "OVERDUE":
                    c.font = Font(name="Calibri", size=10, bold=True, color="FFDC2626")
                    c.fill = PatternFill(start_color="FFFEE2E2", end_color="FFFEE2E2", fill_type="solid")
                elif status_display == "ACTIVE":
                    c.font = Font(name="Calibri", size=10, bold=True, color="FF16A34A")
                    c.fill = PatternFill(start_color="FFDCFCE7", end_color="FFDCFCE7", fill_type="solid")
                elif status_display == "UNDER_MAINTENANCE":
                    c.font = Font(name="Calibri", size=10, bold=True, color="FFD97706")
                    c.fill = PatternFill(start_color="FFFEF3C7", end_color="FFFEF3C7", fill_type="solid")

        current_row += 1

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        # Skip merged title/subtitle when calculating width
        for cell in col[3:]:  # Header and data rows
            if cell.value:
                val_str = str(cell.value)
                max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

    # Set specific generous widths for primary columns
    ws.column_dimensions["A"].width = 10  # Sr. No.
    ws.column_dimensions["B"].width = 16  # Department
    ws.column_dimensions["C"].width = 38  # Machine Name
    ws.column_dimensions["D"].width = 22  # Type
    ws.column_dimensions["E"].width = 22  # Location


def generate_department_excel(department, machines):
    """
    Generates an Excel workbook for the given department and machines.
    Returns BytesIO stream.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    build_department_sheet(ws, department, machines)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_all_departments_excel(department_machines_dict):
    """
    Generates an Excel workbook with a separate tab for each department.
    Returns BytesIO stream.
    """
    wb = openpyxl.Workbook()
    first = True
    for dept, machines in department_machines_dict.items():
        if first:
            ws = wb.active
            first = False
        else:
            ws = wb.create_sheet()
        build_department_sheet(ws, dept, machines)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
